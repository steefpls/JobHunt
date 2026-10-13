"""The candidate profile: contact details (exact, used as-is on every document)
plus a free-text profile (the old info.txt) the model tailors from.
Also reads uploaded CVs (PDF, Word, text) into plain text.
"""
import json
from pathlib import Path

from . import paths

CONTACT_FIELDS = ["name", "email", "phone", "location", "links"]
CV_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}

PROFILE_TEMPLATE = """# CANDIDATE PROFILE

## Summary
(2-3 sentences)

## Work authorization & logistics
Citizenship / PR / visa status, notice period, salary expectation, remote/on-site preference.

## Experience
- Title - Company - Location - Dates
  - Quantified achievement...

## Skills

## Projects

## Education

## Certifications & Languages

## Raw highlights for tailoring
- The 15-25 strongest achievement statements.
"""


def contact_path() -> Path:
    return paths.data_dir() / "contact.json"


def load_contact() -> dict:
    c = {"name": "", "email": "", "phone": "", "location": "", "links": []}
    try:
        c.update(json.loads(contact_path().read_text(encoding="utf-8")))
    except (OSError, json.JSONDecodeError):
        pass
    if isinstance(c.get("links"), str):
        c["links"] = [x.strip() for x in c["links"].splitlines() if x.strip()]
    return c


def save_contact(contact: dict) -> None:
    clean = {k: contact.get(k, "") for k in CONTACT_FIELDS}
    if isinstance(clean["links"], str):
        clean["links"] = [x.strip() for x in clean["links"].splitlines() if x.strip()]
    contact_path().write_text(json.dumps(clean, indent=2), encoding="utf-8")


# Screening answers: the questions application forms ask that a resume doesn't
# answer. The apply agent may only answer these topics from what is saved here;
# a form asking one with no saved answer goes to you as "Needs you" (see
# applyagent.screening_topic). Order = order on the Profile page.
SCREENING_FIELDS = {
    "work_auth": ("Work authorisation", "e.g. Singapore Citizen / PR / Employment Pass holder"),
    "sponsorship": ("Needs visa sponsorship", "Yes or No"),
    "notice_period": ("Notice period / earliest start", "e.g. 1 month"),
    "expected_salary": ("Expected salary", "e.g. SGD 7,500 per month"),
    "current_salary": ("Current salary (leave empty to answer it yourself)", "e.g. SGD 6,800 per month"),
    "years_experience": ("Years of relevant experience", "e.g. 5"),
    "relocate": ("Willing to relocate", "Yes or No"),
}


def screening_path() -> Path:
    return paths.data_dir() / "screening.json"


def load_screening() -> dict:
    s = {k: "" for k in SCREENING_FIELDS}
    try:
        saved = json.loads(screening_path().read_text(encoding="utf-8"))
        s.update({k: str(v).strip() for k, v in saved.items() if k in SCREENING_FIELDS and v})
    except (OSError, json.JSONDecodeError, AttributeError):
        pass
    return s


def save_screening(answers: dict) -> None:
    clean = {k: str(answers.get(k) or "").strip()[:200] for k in SCREENING_FIELDS}
    screening_path().write_text(json.dumps(clean, indent=2), encoding="utf-8")


def load_profile() -> str:
    p = paths.profile_path()
    return p.read_text(encoding="utf-8") if p.exists() else ""


def save_profile(text: str) -> None:
    paths.profile_path().write_text(text.replace("\r\n", "\n"), encoding="utf-8")


def load_scope_text() -> str:
    p = paths.scope_path()
    return p.read_text(encoding="utf-8") if p.exists() else ""


def save_scope_text(text: str) -> None:
    paths.scope_path().write_text(text.replace("\r\n", "\n"), encoding="utf-8")


# The first line of scope.example.md. While it is still in the scope, the scope is
# somebody else's example (salary floor included), so the hunt refuses to run on it.
SCOPE_EXAMPLE_MARKER = "(example: replace with your own)"


def scope_is_example(text: str) -> bool:
    return SCOPE_EXAMPLE_MARKER in (text or "")


def salary_floor() -> float:
    """The expected salary from the screening answers, as SGD a month (0 = none saved)."""
    from .extract import parse_salary_floor
    return parse_salary_floor(load_screening()["expected_salary"])


def list_cvs() -> list[Path]:
    return sorted(p for p in paths.cv_dir().iterdir() if p.suffix.lower() in CV_SUFFIXES)


def cv_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix in (".html", ".htm"):
        from .extract import html_to_text
        return html_to_text(path.read_text(encoding="utf-8", errors="replace"))
    return path.read_text(encoding="utf-8", errors="replace")


def all_cv_texts(max_chars_each: int = 20000) -> list[tuple[str, str]]:
    out = []
    for p in list_cvs():
        try:
            out.append((p.name, cv_text(p)[:max_chars_each]))
        except Exception as e:  # a broken file shouldn't stop the others
            out.append((p.name, f"(could not read this file: {type(e).__name__})"))
    return out


def is_ready() -> bool:
    c = load_contact()
    return bool(c.get("name") and c.get("email") and load_profile().strip())
